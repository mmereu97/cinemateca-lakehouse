"""Prepare the Cinemateca Excel export for ingestion into Databricks.

Why this script exists
----------------------
Spark cannot read .xlsx natively (it needs a third-party library such as
spark-excel, which is not available on Databricks Free Edition serverless
compute). We therefore convert the Excel file locally into **Parquet**:

* Parquet is columnar, compressed and keeps data types — unlike CSV, where
  everything becomes text and commas inside plot summaries or actor lists
  would need fragile quoting/escaping.
* The IMDb link is stored as the *hyperlink* of column H (not as cell text),
  and only openpyxl can read hyperlinks — Spark/pandas `read_excel` cannot.

What it does
------------
1. Loads `Sheet1` with openpyxl (``read_only=False`` because read-only mode
   does not expose hyperlinks).
2. Drops private columns B–E (download links, disk paths, quality, file size)
   so they never leave the local machine.
3. Extracts the IMDb hyperlink of column H into two new columns:
   ``imdb_url`` and ``imdb_id`` (the ``tt\\d+`` identifier).
4. Sanitizes the header (row 1) into snake_case ASCII names; rows 2–3
   (unit and provenance: scrap / manual / algoritm) are saved as a JSON
   metadata sidecar instead of polluting the data.
5. Writes a single Parquet file ready to be uploaded to a Unity Catalog
   Volume, plus ``columns_metadata.json``.

Usage
-----
    python prepare_data.py <input.xlsx> <output_dir>

Example:
    python prepare_data.py Cinemateca.xlsx out/
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

# Columns that must never leave the local machine (private/filesystem info).
DROP_LETTERS = {"B", "C", "D", "E"}

# Column H (1-based index 8) holds the title text + IMDb hyperlink.
TITLE_COL_IDX = 8

HEADER_ROWS = 3  # row 1 = name, row 2 = unit, row 3 = source (scrap/manual/algoritm)

IMDB_ID_RE = re.compile(r"(tt\d+)")


def sanitize_name(raw: str | None, letter: str) -> str:
    """Turn a raw Excel header into a snake_case ASCII identifier.

    Falls back to ``col_<letter>`` when the header cell is empty.
    """
    if raw is None or str(raw).strip() == "":
        return f"col_{letter.lower()}"
    text = unicodedata.normalize("NFKD", str(raw))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^0-9a-zA-Z]+", "_", text).strip("_").lower()
    return text or f"col_{letter.lower()}"


def main(input_path: str, output_dir: str) -> None:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    print(f"Loading {input_path} ...")
    wb = openpyxl.load_workbook(input_path, read_only=False, data_only=True)
    ws = wb["Sheet1"]

    n_cols = ws.max_column
    letters = [get_column_letter(i) for i in range(1, n_cols + 1)]
    keep = [i for i, letter in enumerate(letters, start=1) if letter not in DROP_LETTERS]

    # --- Build sanitized, de-duplicated column names from header row 1 ---
    names: list[str] = []
    seen: set[str] = set()
    meta: dict[str, dict] = {}
    for i in keep:
        letter = letters[i - 1]
        name = sanitize_name(ws.cell(row=1, column=i).value, letter)
        if name in seen:  # e.g. "gen" appears twice -> gen, gen_l
            name = f"{name}_{letter.lower()}"
        seen.add(name)
        names.append(name)
        meta[name] = {
            "excel_column": letter,
            "original_name": ws.cell(row=1, column=i).value,
            "unit": ws.cell(row=2, column=i).value,
            "source": ws.cell(row=3, column=i).value,
        }

    # --- Read data rows, extracting the IMDb hyperlink from column H ---
    records: list[list] = []
    urls: list[str | None] = []
    for row in range(HEADER_ROWS + 1, ws.max_row + 1):
        title_cell = ws.cell(row=row, column=TITLE_COL_IDX)
        # Skip fully empty rows (no title and no movie_id).
        if title_cell.value is None and ws.cell(row=row, column=1).value is None:
            continue
        records.append([ws.cell(row=row, column=i).value for i in keep])
        urls.append(title_cell.hyperlink.target if title_cell.hyperlink else None)

    df = pd.DataFrame(records, columns=names)
    df["imdb_url"] = urls
    df["imdb_id"] = df["imdb_url"].str.extract(IMDB_ID_RE, expand=False)

    # --- Normalize dtypes so Parquet writing never hits mixed-type columns.
    # Bronze stays "raw": numeric columns stay numeric when every value parses,
    # everything else becomes a plain string. Proper typing happens in Silver.
    for col in df.columns:
        if df[col].dtype == object:
            as_num = pd.to_numeric(df[col], errors="coerce")
            non_null = df[col].notna()
            if non_null.any() and as_num[non_null].notna().all():
                df[col] = as_num
            else:
                df[col] = df[col].astype("string")

    parquet_path = out / "cinemateca_movies.parquet"
    df.to_parquet(parquet_path, index=False)
    (out / "columns_metadata.json").write_text(
        json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    print(f"Wrote {parquet_path}  ({len(df)} movies, {len(df.columns)} columns)")
    print(f"Wrote {out / 'columns_metadata.json'}")
    print(f"Movies with an IMDb link: {df['imdb_url'].notna().sum()}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Usage: python prepare_data.py <input.xlsx> <output_dir>")
    main(sys.argv[1], sys.argv[2])
